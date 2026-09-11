from pathlib import Path
from pypdf import PdfReader
from PIL import Image, ImageOps, ImageDraw, ImageFont
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.lib.units import mm
from openpyxl import load_workbook
import json, zipfile

base=Path('C:/Projects/zautomatyzujemy')
out=base/'output/sprzedaz-zautomatyzujemy'
work=base/'tmp/sales-kit'
report=[]
expected={'01-podrecznik-sprzedazy.pdf':29,'02-ulotki-A5.pdf':6,'02-ulotki-A5-druk.pdf':6,'03-wizytowka-85x55.pdf':2,'03-wizytowka-85x55-druk.pdf':2,'04-sciaga-na-spotkanie-A4.pdf':2,'06-karty-rozmowy-i-odbioru-A4.pdf':2}
for name,pages in expected.items():
    reader=PdfReader(out/name)
    assert len(reader.pages)==pages
    for i,p in enumerate(reader.pages):
        txt=p.extract_text()
        assert len(txt)>30,(name,i)
        assert '\ufffd' not in txt,(name,i)
        if name.startswith(('02','03')):
            assert 'zautomatyzujemy.pl' in txt
        if '-druk' in name:
            w,h=(148,210) if name.startswith('02') else (85,55)
            assert abs(float(p.trimbox.width)/mm-w)<.01
            assert abs(float(p.trimbox.height)/mm-h)<.01
            assert abs(float(p.bleedbox.width)/mm-(w+6))<.01
            assert b' k' in p.get_contents().get_data()
        for f in p['/Resources']['/Font'].get_object().values():
            font=f.get_object()
            if '/FontDescriptor' in font:
                assert any(k in font['/FontDescriptor'].get_object() for k in ['/FontFile','/FontFile2','/FontFile3'])
    report.append({'pdf':name,'pages':pages,'dimensions_and_text':'OK'})

# Sample every QR module from final print raster, including its quiet zone.
# This verifies the rendered matrix against the encoded URL without claiming a phone scan.
for filename,url,w,h,x,y,size in [
 ('flyer-1.png','https://zautomatyzujemy.pl/#kontakt',160*mm,222*mm,6*mm+148*mm-100,6*mm+482,71),
 ('flyer-3.png','https://zautomatyzujemy.pl/#kontakt',160*mm,222*mm,6*mm+148*mm-100,6*mm+482,71),
 ('flyer-5.png','https://zautomatyzujemy.pl/#kontakt',160*mm,222*mm,6*mm+148*mm-100,6*mm+482,71),
 ('card-2.png','https://zautomatyzujemy.pl/#klara',97*mm,67*mm,6*mm+13,6*mm+96,50)]:
    widget=QrCodeWidget(url,barLevel='M');widget.qr.make();matrix=widget.qr.modules
    n=len(matrix);im=Image.open(work/filename).convert('L');mismatch=0
    for r in range(n+8):
        for c in range(n+8):
            px=int((x+(c+.5)*size/(n+8))*im.width/w)
            py=int((y+(r+.5)*size/(n+8))*im.height/h)
            expected_dark=bool(matrix[r-4][c-4]) if 4<=r<n+4 and 4<=c<n+4 else False
            actual=im.getpixel((px,py))<128
            mismatch+=actual!=expected_dark
    assert mismatch==0,(filename,mismatch)
    report.append({'qr_raster':filename,'encoded_url':url,'mismatched_modules':mismatch})

formulas=load_workbook(out/'05-narzedzia-sprzedazy.xlsx',data_only=False)
cache=load_workbook(out/'05-narzedzia-sprzedazy.xlsx',data_only=True)
assert len(cache.sheetnames)==6
count=0
for sheet in formulas:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f':
                count+=1
                cached=cache[sheet.title][cell.coordinate]
                assert cached.value is not None,(sheet.title,cell.coordinate)
                assert cached.data_type!='e',(sheet.title,cell.coordinate,cached.value)
assert cache['Kalkulator korzyści']['C5'].value==28
assert cache['Kalkulator korzyści']['C7'].value==-550
assert cache['Kalkulator korzyści']['C9'].value=='Brak zwrotu'
assert all(cache['Lejek CRM'].cell(r,2).value is None for r in range(5,35))
report.append({'workbook':'05-narzedzia-sprzedazy.xlsx','sheets':len(cache.sheetnames),'cached_formulas':count,'crm_empty_records':30})

# Presentation preview from the delivered design, not a mock client implementation.
preview=Image.new('RGB',(1800,1350),'#e8e4dd')
draw=ImageDraw.Draw(preview)
title_font=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',46)
small_font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
draw.text((65,38),'Zautomatyzujemy.pl',font=title_font,fill='#151719')
draw.text((65,99),'Zestaw do rozmów, ofert i wdrożeń',font=small_font,fill='#62625d')
for file,x in [('book-01.png',65),('flyer-1.png',495),('flyer-3.png',925),('flyer-5.png',1355)]:
    im=ImageOps.contain(Image.open(work/file).convert('RGB'),(385,570))
    preview.paste(im,(x,168))
for file,x in [('card-1.png',65),('card-2.png',660)]:
    im=ImageOps.contain(Image.open(work/file).convert('RGB'),(560,430))
    preview.paste(im,(x,805))
draw.text((1270,842),'29 stron podręcznika',font=small_font,fill='#151719')
draw.text((1270,890),'3 dwustronne ulotki',font=small_font,fill='#151719')
draw.text((1270,938),'Wizytówka i formularze',font=small_font,fill='#151719')
draw.text((1270,986),'Kalkulator i rejestr klientów',font=small_font,fill='#151719')
draw.text((1270,1034),'Szablon oferty i wiadomości',font=small_font,fill='#151719')
preview.save(out/'podglad-zestawu.jpg',quality=94)

(work/'verification-final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
archive=base/'output/zautomatyzujemy-pakiet-sprzedazowy.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
    for file in sorted(out.rglob('*')):
        if file.is_file():z.write(file,Path('zautomatyzujemy-pakiet-sprzedazowy')/file.relative_to(out))
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    report.append({'archive':str(archive),'files':len(z.namelist()),'bytes':archive.stat().st_size})
print(json.dumps(report,ensure_ascii=False,indent=2))
