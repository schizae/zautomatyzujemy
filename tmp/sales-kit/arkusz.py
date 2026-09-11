from pathlib import Path
import xlsxwriter
from openpyxl import load_workbook

OUT=Path('C:/Projects/zautomatyzujemy/output/sprzedaz-zautomatyzujemy/05-narzedzia-sprzedazy.xlsx')
wb=xlsxwriter.Workbook(OUT)
wb.set_properties({'title':'Narzędzia sprzedaży Zautomatyzujemy.pl','author':'Zautomatyzujemy.pl'})
wb.set_calc_mode('auto')
BG='#f5f2ed'; INK='#151719'; ACC='#c93820'; BLUE='#e7f1fb'
def fmt(**kw): return wb.add_format({'font_name':'Arial','font_size':11,'font_color':INK,'valign':'vcenter','bg_color':BG,**kw})
body=fmt(text_wrap=True); note=fmt(text_wrap=True,font_color='#5b6065'); title=fmt(font_size=17,bold=True); head=fmt(bold=True,bg_color=INK,font_color='white',text_wrap=True)
entry=fmt(bg_color=BLUE,text_wrap=True,font_color='#14528f'); num=fmt(num_format='#,##0.0;(#,##0.0);0.0'); money=fmt(num_format='#,##0" zł";(#,##0)" zł";0" zł"')
pct=fmt(num_format='0.0%;(0.0%);0.0%'); inpnum=fmt(bg_color=BLUE,font_color='#14528f',num_format='#,##0.0'); inppct=fmt(bg_color=BLUE,font_color='#14528f',num_format='0.0%'); date=fmt(bg_color=BLUE,font_color='#14528f',num_format='dd.mm.yyyy'); bad=fmt(bg_color='#fbe0db',font_color=ACC,bold=True)
def sheet(name, subtitle, rows=42):
    s=wb.add_worksheet(name); s.hide_gridlines(2); s.set_tab_color(ACC if name=='START' else INK)
    s.set_column('A:A',3); s.set_column('B:B',49); s.set_column('C:C',25); s.set_column('D:D',65)
    for r in range(rows):
        s.set_row(r,30)
        for c in range(1,4): s.write_blank(r,c,None,body)
    s.set_row(1,32); s.write('B2',name,title); s.write('B3',subtitle,fmt(font_color='#5b6065')); s.set_row(2,46)
    s.set_landscape(); s.set_paper(9); s.fit_to_pages(1,0); s.set_margins(.3,.3,.4,.4); s.print_area(0,1,rows-1,3)
    s.set_footer('&LZautomatyzujemy.pl&R&N / &P'); s.set_zoom(90)
    return s
def line(s,r,label,value=None,explain='',style=None):
    s.write(r-1,1,label,body)
    if value is None: s.write_blank(r-1,2,None,entry)
    elif isinstance(value,str) and value.startswith('='): s.write_formula(r-1,2,value,style or num)
    else: s.write(r-1,2,value,style or entry)
    s.write(r-1,3,explain,note)
def formula(s,cell,text,cached,style=num): s.write_formula(cell,text,style,cached)
def section(s,r,text):
    s.write(r-1,1,text,head); s.write_blank(r-1,2,None,head); s.write_blank(r-1,3,None,head)

s=sheet('START','Zautomatyzujemy.pl — roboczy zestaw do rozmowy, wyceny i prowadzenia wdrożenia.',29)
instructions=[
('1. Discovery','Zapisz proces, jego skalę, wyjątki i właściciela. Poproś o zanonimizowaną próbkę. Ustal źródła liczb.'),
('2. Kalkulator korzyści','Zastąp wszystkie HIPOTETYCZNE liczby danymi klienta. Błękitne komórki są do edycji. Komórki bez błękitu zawierają wyniki lub instrukcje.'),
('3. Wycena','Oszacuj pracę po potwierdzeniu zakresu. Podatek jest opcjonalnym wejściem. Kalkulator korzyści i wycena są niezależne: przenieś uzgodnione koszty ręcznie.'),
('4. Lejek CRM','Wpisuj rzeczywiste firmy. Każda otwarta szansa ma następny krok i datę. Zapisz źródło relacji oraz podstawę kontaktu, jeśli dotyczy.'),
('5. Odbiór','Uzgodnij kryteria przed budową. Przy odbiorze dodaj dowody i decyzję osoby uprawnionej. Samo zaznaczenie listy nie zastępuje akceptacji klienta.'),
('Dane i założenia','Wszystkie liczby początkowe w kalkulatorze i wycenie są HIPOTETYCZNE, nie są cennikiem ani obietnicą wyniku. Zero to świadome założenie; puste wymagane pole wstrzymuje wynik.'),
('Czas a pieniądze','Uwolnione godziny oznaczają dostępną pojemność zespołu. Korzyść gotówkowa wymaga faktycznie unikniętego wydatku, np. nadgodzin. Udział zamieniony na gotówkę domyślnie wynosi 0%.'),
('Dodatkowa marża','Osobny wariant uwzględnia marżę po kosztach realizacji. Nie wpisuj przychodu. Nie doliczaj ponownie wartości godzin lub redukcji kosztów ujętych w bazie.'),
('Horyzont i kwoty','Model zakłada stały miesięczny wolumen i korzyści od pierwszego miesiąca po uruchomieniu. Nie modeluje rozruchu, dyskonta ani opóźnień płatności. Koszty podawaj na tej samej podstawie, z nieodliczalnymi podatkami, jeżeli występują.'),
('Zapisywanie','Zapisz kopię pliku dla konkretnego klienta. Po zmianach otwórz w Excelu lub zgodnym arkuszu z automatycznym przeliczaniem. Nie nadpisuj formuł.')]
for r,(label,txt) in enumerate(instructions,5):
    s.write(r-1,1,label,body); s.write(r-1,3,txt,note); s.set_row(r-1,66)
s.print_area('B1:D15')

s=sheet('Kalkulator korzyści','HIPOTETYCZNY przykład. PLN. Wypełnij C15:C26; zero jest wartością, puste pole zatrzymuje wynik.',44)
section(s,4,'Wynik do rozmowy')
for r,l in [(5,'Uwolniony czas (godz./mies.)'),(6,'Wartość czasu (zł/mies.)'),(7,'Korzyść gotówkowa netto (zł/mies.)'),(8,'ROI gotówkowe w horyzoncie'),(9,'Zwrot wdrożenia (miesiące)'),(10,'Wynik z dodatkową marżą (zł/mies.)'),(11,'ROI z dodatkową marżą')]: line(s,r,l,0,style=num)
section(s,14,'Założenia — wszystkie liczby HIPOTETYCZNE')
inputs=[('Zadania / miesiąc',300,'Powtarzalne wykonania tego samego procesu.'),('Minuty przed / zadanie',10,'Średnia z pomiaru, wraz z obsługą wyjątków.'),('Minuty po / zadanie',2,'Uwzględnij kontrolę człowieka i utrzymujące się wyjątki.'),('Udział zadań automatyzowanych',.7,'0–100%. Pozostałe zadania nie zmieniają czasu.'),('Pełny koszt godziny (zł)',70,'Koszt pracodawcy lub realny koszt zewnętrzny.'),('Udział czasu zamieniony na gotówkę',0,'0–100%. Bez potwierdzonej redukcji wydatków zostaw 0%.'),('Koszt wdrożenia (zł)',8000,'Jednorazowy nakład, wraz z pracą klienta, jeśli wyceniona.'),('Opieka miesięczna (zł)',400,'Koszt stały utrzymania.'),('Narzędzia miesięcznie (zł)',150,'Licencje, API i infrastruktura.'),('Horyzont (miesiące)',12,'Dodatnia liczba całkowita po uruchomieniu.'),('Dodatkowa marża miesięczna (zł)',0,'Po kosztach realizacji, bez korzyści gotówkowej ujętej wyżej.'),('Realizacja dodatkowej marży',.5,'Ostrożny udział 0–100%. Dotyczy tylko dodatkowej marży.')]
for r,(l,v,n) in enumerate(inputs,15):
    line(s,r,l,v,n,inppct if r in [18,20,26] else inpnum)
    s.data_validation(r-1,2,r-1,2,{'validate':'decimal','criteria':'between','minimum':0 if r!=24 else 1,'maximum':1 if r in [18,20,26] else 100000000,'error_type':'stop','error_message':'Podaj liczbę w dopuszczalnym zakresie.','ignore_blank':False})
s.data_validation('C24',{'validate':'integer','criteria':'between','minimum':1,'maximum':1200,'ignore_blank':False})
section(s,28,'Obliczenia i interpretacja')
line(s,29,'Stan założeń',0,'Jeśli wklejasz dane, sprawdź również zakresy procentowe. Walidacja wklejania zależy od aplikacji.')
formula(s,'C29','=IF(COUNT(C15:C26)<>12,"Uzupełnij liczby",IF(OR(MIN(C15:C26)<0,C18>1,C20>1,C26>1,C24<1,MOD(C24,1)<>0),"Popraw zakresy","OK"))','OK',body)
calc={30:('Uwolnione godziny / miesiąc','C15*(C16-C17)*C18/60',28),31:('Wartość uwolnionego czasu','C30*C19',1960),32:('Redukcja wydatków / miesiąc','C31*C20',0),33:('Koszty stałe / miesiąc','SUM(C22:C23)',550),34:('Korzyść gotówkowa netto / miesiąc','C32-C33',-550),35:('Całkowity koszt w horyzoncie','C21+C33*C24',14600),36:('Wynik gotówkowy w horyzoncie','C34*C24-C21',-14600),37:('Marża dodatkowa po korekcie','C25*C26',0),38:('Wynik z marżą / miesiąc','C34+C37',-550),39:('Wynik z marżą w horyzoncie','C38*C24-C21',-14600)}
for r,(l,f,c) in calc.items():
    line(s,r,l,0); formula(s,f'C{r}',f'=IF(C29<>"OK","n.d.",{f})',c,num if r==30 else money)
for r,ref,c,style in [(5,30,28,num),(6,31,1960,money),(7,34,-550,money),(10,38,-550,money)]: formula(s,f'C{r}',f'=C{ref}',c,style)
formula(s,'C8','=IF(C29<>"OK","n.d.",IF(C35=0,"n.d.",C36/C35))',-1,pct)
formula(s,'C9','=IF(C29<>"OK","n.d.",IF(C34<=0,"Brak zwrotu",IF(C21=0,0,C21/C34)))','Brak zwrotu',num)
formula(s,'C11','=IF(C29<>"OK","n.d.",IF(C35=0,"n.d.",C39/C35))',-1,pct)
line(s,41,'Próg: redukcja wydatków / miesiąc',0,'Wymagana redukcja wydatków, aby odzyskać wdrożenie w zadanym horyzoncie. Wariant bazowy.')
formula(s,'C41','=IF(C29<>"OK","n.d.",C33+C21/C24)',1216.666666667,money)
s.write('D6','Wycena pojemności zespołu. Nie dodawaj do wyniku gotówkowego.',note)
s.write('D8','(Redukcja wydatków minus wszystkie koszty) / wszystkie koszty. Koszt = 0 oznacza n.d.',note)
s.write('D9','Stałe przepływy od uruchomienia. Zwrot może wypaść poza horyzontem. Przy wyniku ≤ 0: brak zwrotu.',note)
s.set_row(8,46); s.set_row(7,46)
s.conditional_format('C7:C11',{'type':'cell','criteria':'<','value':0,'format':bad}); s.freeze_panes(14,2)

s=sheet('Wycena','HIPOTETYCZNY przykład. Uzupełnij własne godziny, stawkę, koszty i opcjonalną stawkę podatku.',27)
section(s,4,'Nakład pracy i założenia')
items=[('Analiza (godz.)',6),('Budowa (godz.)',24),('Testy (godz.)',8),('Szkolenie i przekazanie (godz.)',2),('Bufor czasu',.2),('Stawka sprzedażowa (zł/godz.)',200),('Koszty jednorazowe poza pracą (zł)',0),('Podatek doliczany do ceny — opcjonalny',None)]
for r,(l,v) in enumerate(items,5):
    line(s,r,l,v,'Pozostaw puste, gdy nie dotyczy. To nie jest założenie stawki VAT.' if r==12 else '',inppct if r in [9,12] else inpnum)
    s.data_validation(r-1,2,r-1,2,{'validate':'decimal','criteria':'between','minimum':0,'maximum':1 if r in [9,12] else 100000000})
section(s,14,'Kalkulacja ceny')
formula(s,'C15','=IF(COUNT(C5:C11)<>7,"Uzupełnij liczby",IF(OR(MIN(C5:C11)<0,C9>1,AND(C12<>"",NOT(ISNUMBER(C12))),C12<0,C12>1),"Popraw zakresy","OK"))','OK',body); s.write('B15','Stan założeń',body)
for r,l,f,c,sty in [(16,'Godziny przed buforem','SUM(C5:C8)',40,num),(17,'Godziny z buforem','C16*(1+C9)',48,num),(18,'Cena przed podatkiem (zł)','C17*C10+C11',9600,money),(19,'Podatek doliczony (zł)','IF(C12="",0,C18*C12)',0,money),(20,'Cena z zadanym podatkiem (zł)','SUM(C18:C19)',9600,money)]:
    line(s,r,l,0); formula(s,f'C{r}',f'=IF(C15<>"OK","n.d.",{f})',c,sty)
line(s,22,'Opieka / miesiąc (zł)',None,'Uzgodnij zakres, limit godzin, czas reakcji i prace dodatkowe.',inpnum)
line(s,23,'Narzędzia / miesiąc (zł)',None,'Wskaż kto płaci licencje i opłaty za zużycie.',inpnum)
line(s,24,'Zakres i wyłączenia',None,'Dopisz konkretne procesy, integracje, limity i wyłączenia.')
s.set_row(23,70)

s=sheet('Lejek CRM','Pusty rejestr 30 szans. Dodawaj wyłącznie rzeczywiste kontakty i ustalony następny krok.',35)
headers=['Firma','Branża','Osoba / rola','Źródło relacji / zgody','Problem','Wolumen / miesiąc','Etap','Wartość (zł)','Następny krok','Data kroku','Wynik / powód']
widths=[25,20,25,36,40,22,24,19,40,17,35]
for c,(h,w) in enumerate(zip(headers,widths),1):
    s.set_column(c,c,w); s.write(3,c,h,head)
    for r in range(4,34): s.write_blank(r,c,None,date if c==10 else inpnum if c==8 else entry)
s.set_row(3,42); s.freeze_panes(4,2); s.autofilter(3,1,33,11)
s.data_validation('H5:H34',{'validate':'list','source':['Do kontaktu','Rozmowa umówiona','Discovery','Oferta','Negocjacje','Wygrana','Przegrana','Wstrzymana']})
s.conditional_format('B5:L34',{'type':'formula','criteria':'AND($B5<>"",$K5<>"",$K5<TODAY(),$H5<>"Wygrana",$H5<>"Przegrana")','format':bad})
s.print_area('B1:L34'); s.set_paper(8); s.fit_to_pages(1,2)

s=sheet('Discovery','Uzupełnij podczas rozmowy. Ustal pomiar bazowy przed obietnicą efektu.',27)
s.set_column('B:B',40); s.set_column('C:C',45); s.set_column('D:D',54)
questions=[('Firma i rozmówca','Nazwa, rola, data rozmowy.'),('Właściciel procesu','Kto odpowiada za wynik i akceptuje zmianę?'),('Proces do usprawnienia','Co go uruchamia, jakie ma wejście i oczekiwany wynik?'),('Problem i konsekwencja','Co boli i jaki ma wpływ na firmę?'),('Wolumen i sezonowość','Ile przypadków miesięcznie, w szczycie i w miesiącu spokojnym?'),('Czas bazowy','Pomiar na ilu przypadkach? Średnia, rozrzut, źródło pomiaru.'),('Wyjątki i błędy','Jak często wymagają ręcznej obsługi? Kto ją wykonuje?'),('Systemy i dostęp','Aplikacje, API, dokumentacja, testowe konto i właściciel dostępu.'),('Dane i poufność','Jakie dane są konieczne, gdzie mogą trafić, czas przechowywania?'),('Próbka do analizy','Uzgodnij zanonimizowany przykład wejścia i poprawnego wyniku.'),('Cel i pomiar sukcesu','Miernik, wartość bazowa, cel, metoda i termin pomiaru.'),('Ryzyka i ograniczenia','Reguły biznesowe, ręczna kontrola, cofnięcie zmiany.'),('Zakres pilota','Jeden proces, systemy, wyjątki, limit wolumenu, wyłączenia.'),('Decyzja i budżet','Kto decyduje, jaka ścieżka zakupu, budżet i termin?'),('Odbiór i odpowiedzialności','Kto testuje, akceptuje i utrzymuje proces?'),('Następny krok','Konkretna czynność, właściciel i data.')]
for r,(l,n) in enumerate(questions,5): line(s,r,l,None,n); s.set_row(r-1,62)
s.freeze_panes(4,2); s.print_area('B1:D20')

s=sheet('Odbiór','Uzgodnij kryteria przed pracą. Każda pozycja wymaga wyniku testu lub uzasadnienia „Nie dotyczy”.',26)
s.set_column('B:B',45); s.set_column('C:C',20); s.set_column('D:D',58); s.set_column('E:E',24)
for c,h in enumerate(['Kryterium','Status','Dowód / uwagi / uzasadnienie','Właściciel / data'],1): s.write(3,c,h,head)
checks=['Zakres zgodny z uzgodnieniem','Scenariusz standardowy przeszedł test','Uzgodnione wyjątki obsłużone','Błędne dane nie psują procesu','Powtórzenie nie tworzy duplikatów','Alert błędu trafia do właściciela','Ręczne przejęcie procesu sprawdzone','Dostępy i sekrety przekazane bezpiecznie','Logi i retencja zgodne z ustaleniem','Pomiar celu wykonany i opisany','Instrukcja operacyjna przekazana','Osoba klienta przeszkolona','Opieka i odpowiedzialności ustalone','Otwarta lista usterek uzgodniona']
for r,l in enumerate(checks,4):
    s.write(r,1,l,body)
    for c in [2,3,4]: s.write_blank(r,c,None,entry)
    s.set_row(r,46)
s.data_validation('C5:C18',{'validate':'list','source':['Do sprawdzenia','Spełnione','Niespełnione','Nie dotyczy']})
s.conditional_format('C5:C18',{'type':'text','criteria':'containing','value':'Niespełnione','format':bad})
for r,l in [(20,'Decyzja odbioru'),(21,'Akceptujący po stronie klienta'),(22,'Data decyzji'),(23,'Warunki / pozostałe działania')]: line(s,r,l,None)
s.data_validation('C20',{'validate':'list','source':['Przyjęte','Przyjęte warunkowo','Nieprzyjęte']}); s.set_row(22,64)
s.print_area('B1:E23'); s.freeze_panes(4,2)
wb.close()

book=load_workbook(OUT,data_only=False); cached=load_workbook(OUT,data_only=True)
assert len(book.sheetnames)==6
assert all(book['Lejek CRM'].cell(r,c).value is None for r in range(5,35) for c in range(2,13))
assert cached['Kalkulator korzyści']['C8'].value == -1
assert cached['Wycena']['C20'].value == 9600
assert all(c.data_type!='e' for sh in cached for row in sh for c in row)
assert all(c.value is not None for sh in book for row in sh for c in row if c.data_type=='f')
print('Export OK:',OUT,'; 6 sheets; 30 empty CRM rows; cached ROI -100%, quote 9600 PLN; no cached errors')
